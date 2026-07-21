"""덱 HTML → PPTX 변환 (기본: vibe-coding/ch01-setup/1-understand.html).
- 인자로 slide/ 기준 상대경로 지정 가능: python export_pptx.py vibe-coding/ch10-desktop/1-desktop-program.html
- Chromium headless로 4:3 viewport 캡처 (1920x1440), .slide 개수 자동 감지
- 각 #slide=N 마다 PNG → python-pptx로 4:3 PPTX에 풀-블리드 삽입
"""
import asyncio
from pathlib import Path
from playwright.async_api import async_playwright
from pptx import Presentation
from pptx.util import Inches

import sys

ROOT = Path(__file__).resolve().parent
# 변환할 덱 — 인자로 slide/ 기준 상대경로 지정 가능 (기본: 1장 이해)
_rel = sys.argv[1] if len(sys.argv) > 1 else 'vibe-coding/ch01-setup/1-understand.html'
HTML_PATH = Path(_rel) if Path(_rel).is_absolute() else (ROOT / 'slide' / _rel)
DIST = ROOT / 'dist'
PNG_DIR = DIST / f'{HTML_PATH.stem}-png'
PPTX_PATH = DIST / f'{HTML_PATH.stem}.pptx'

VIEW_W, VIEW_H = 1920, 1440


async def capture():
    DIST.mkdir(exist_ok=True)
    PNG_DIR.mkdir(exist_ok=True)
    images = []
    async with async_playwright() as p:
        browser = await p.chromium.launch()
        context = await browser.new_context(
            viewport={'width': VIEW_W, 'height': VIEW_H},
            device_scale_factor=1,
        )
        page = await context.new_page()
        base_url = HTML_PATH.as_uri()
        await page.goto(base_url)
        await page.wait_for_load_state('networkidle')
        n = await page.evaluate("document.querySelectorAll('.slide').length")
        for i in range(1, n + 1):
            await page.goto(f'{base_url}#slide={i}')
            await page.wait_for_load_state('networkidle')
            await page.wait_for_timeout(400)
            out = PNG_DIR / f'slide-{i:02d}.png'
            await page.screenshot(path=str(out), full_page=False)
            images.append(out)
            print(f'  [{i:02d}/{n}] captured')
        await browser.close()
    return images


def build_pptx(images):
    prs = Presentation()
    prs.slide_width = Inches(10)      # 4:3
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for img in images:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(img), 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(PPTX_PATH)
    print(f'\nsaved: {PPTX_PATH}')


if __name__ == '__main__':
    print(f'capturing slides from {HTML_PATH.name}')
    imgs = asyncio.run(capture())
    print('\nbuilding pptx...')
    build_pptx(imgs)
